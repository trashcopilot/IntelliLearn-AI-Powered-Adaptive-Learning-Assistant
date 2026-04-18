import logging
import io
import os
import re
import tempfile

from functools import lru_cache

logger = logging.getLogger(__name__)

try:
    import pdfplumber
except Exception:
    pdfplumber = None

try:
    from google import genai
    from google.genai import types
except Exception:
    genai = None
    types = None


_PDF_TEXT_MIN_CHARS = int(os.getenv('PDF_TEXT_MIN_CHARS', '120'))
_PDF_TEXT_MIN_WORDS = int(os.getenv('PDF_TEXT_MIN_WORDS', '20'))
_PDF_OCR_MAX_PAGES = int(os.getenv('PDF_OCR_MAX_PAGES', '8'))
_PDF_OCR_RENDER_DPI = int(os.getenv('PDF_OCR_RENDER_DPI', '160'))
_PDF_OCR_MAX_OUTPUT_TOKENS = int(os.getenv('PDF_OCR_MAX_OUTPUT_TOKENS', '1500'))
_PDF_OCR_MODEL = os.getenv('GEMINI_OCR_MODEL', os.getenv('GEMINI_MODEL', 'gemini-3-flash-preview'))


def _get_gemini_api_key() -> str:
    for env_name in ('GEMINI_API_KEY', 'GOOGLE_API_KEY', 'GOOGLE_GENAI_API_KEY'):
        api_key = os.getenv(env_name)
        if api_key:
            return api_key
    return ''


@lru_cache(maxsize=1)
def _get_gemini_client():
    api_key = _get_gemini_api_key()
    if not api_key or genai is None:
        return None

    try:
        return genai.Client(api_key=api_key)
    except Exception as exc:
        logger.warning('Gemini OCR client initialization failed: %s', exc)
        return None


def _extract_pdf(path: str) -> str:
    if pdfplumber is None:
        logger.warning('pdfplumber package is unavailable for PDF extraction.')
        return ''

    page_texts = {}
    sparse_pages = []

    try:
        with pdfplumber.open(path) as pdf:
            total_pages = len(pdf.pages)
            for index, page in enumerate(pdf.pages, start=1):
                raw_text = page.extract_text() or ''
                cleaned_text = _clean_extracted_text(raw_text)
                if _page_needs_ocr(cleaned_text):
                    sparse_pages.append((index, page, total_pages))
                if cleaned_text:
                    page_texts[index] = cleaned_text

        if sparse_pages:
            ocr_page_limit = min(len(sparse_pages), _PDF_OCR_MAX_PAGES)
            for index, page, total_pages in sparse_pages[:ocr_page_limit]:
                ocr_text = _ocr_pdf_page_with_gemini(page, index, total_pages)
                if ocr_text:
                    page_texts[index] = ocr_text
    except Exception as exc:
        logger.error('PDF extraction failed for %s: %s', path, exc)
        return ''

    ordered_pages = [f'Page {index}\n{text}' for index, text in sorted(page_texts.items()) if text]
    return '\n\n'.join(ordered_pages)


def _extract_docx(path: str) -> str:
    from docx import Document
    doc = Document(path)
    return '\n'.join(para.text for para in doc.paragraphs)


def _extract_txt(path: str) -> str:
    with open(path, 'r', encoding='utf-8', errors='ignore') as handle:
        return handle.read()


def _extract_audio(path: str) -> str:
    try:
        import speech_recognition as sr
    except Exception:
        logger.warning('speech_recognition package is unavailable for audio transcription.')
        return ''

    recognizer = sr.Recognizer()
    with sr.AudioFile(path) as source:
        audio_data = recognizer.record(source)

    try:
        return recognizer.recognize_google(audio_data)
    except Exception as exc:
        logger.warning('Audio transcription failed for %s: %s', path, exc)
        return ''


def _extract_video(path: str) -> str:
    try:
        from moviepy import VideoFileClip
    except Exception:
        logger.warning('moviepy package is unavailable for video transcription.')
        return ''

    temp_audio = None
    clip = None
    try:
        clip = VideoFileClip(path)
        if clip.audio is None:
            return ''
        with tempfile.NamedTemporaryFile(suffix='.wav', delete=False) as temp_file:
            temp_audio = temp_file.name
        clip.audio.write_audiofile(temp_audio, logger=None)
        return _extract_audio(temp_audio)
    except Exception as exc:
        logger.warning('Video transcription failed for %s: %s', path, exc)
        return ''
    finally:
        if clip is not None:
            clip.close()
        if temp_audio and os.path.exists(temp_audio):
            os.remove(temp_audio)


def _clean_extracted_text(text: str) -> str:
    """Remove formatting artifacts, placeholders, and symbols from extracted text."""
    if not text:
        return text

    placeholder_patterns = (
        r'You can enter a subtitle here.*?(?=$)',
        r'click to edit.*?(?=$)',
        r'add title.*?(?=$)',
    )

    cleaned_lines = []
    for raw_line in text.splitlines():
        line = (raw_line or '').strip()
        if not line:
            continue

        # Remove common slide placeholder lines.
        skip_line = False
        for pattern in placeholder_patterns:
            if re.search(pattern, line, flags=re.IGNORECASE):
                skip_line = True
                break
        if skip_line:
            continue

        line = re.sub(r'[→↓↑←↔]', '', line)

        # Remove pure page numbers like "12" or "12/44".
        if re.fullmatch(r'\d{1,3}(?:\s*/\s*\d{1,3})?', line):
            continue

        # Normalize bullet markers but keep list structure.
        line = re.sub(r'^\s*[•\-\*o]\s*', '- ', line)
        line = re.sub(r'^\s*(\d{1,3})[\.)]\s*', r'\1. ', line)

        # Collapse duplicate neighboring words while preserving sentence structure.
        line = re.sub(r'\b(\w+)\s+(?=\1\b)', '', line, flags=re.IGNORECASE)
        line = re.sub(r'\s+', ' ', line).strip()

        if line:
            cleaned_lines.append(line)

    if not cleaned_lines:
        return ''

    # Preserve line boundaries so downstream summarization can detect headings/lists.
    return '\n'.join(cleaned_lines)


def _page_needs_ocr(text: str) -> bool:
    cleaned = (text or '').strip()
    if not cleaned:
        return True

    words = cleaned.split()
    if len(cleaned) < _PDF_TEXT_MIN_CHARS:
        return True
    if len(words) < _PDF_TEXT_MIN_WORDS:
        return True
    return False


def _render_pdf_page_to_png(page) -> bytes:
    try:
        rendered = page.to_image(resolution=_PDF_OCR_RENDER_DPI)
        image = rendered.original
        if image is None:
            return b''

        buffer = io.BytesIO()
        image.save(buffer, format='PNG')
        return buffer.getvalue()
    except Exception as exc:
        logger.warning('PDF page rendering failed for OCR: %s', exc)
        return b''


def _ocr_pdf_page_with_gemini(page, page_number: int, total_pages: int) -> str:
    client = _get_gemini_client()
    if client is None or types is None:
        return ''

    image_bytes = _render_pdf_page_to_png(page)
    if not image_bytes:
        return ''

    prompt = (
        'You are reading a PDF page image. Extract the visible text faithfully in reading order. '
        'Preserve headings, bullets, numbered lists, and table rows where possible. '
        'Do not summarize, explain, or add commentary. Return only the text you can read.\n\n'
        f'Page {page_number} of {total_pages}.'
    )

    try:
        response = client.models.generate_content(
            model=_PDF_OCR_MODEL,
            contents=[
                prompt,
                types.Part.from_bytes(data=image_bytes, mime_type='image/png'),
            ],
            config={
                'temperature': 0.0,
                'max_output_tokens': _PDF_OCR_MAX_OUTPUT_TOKENS,
            },
        )
        return _clean_extracted_text((getattr(response, 'text', '') or '').strip())
    except Exception as exc:
        logger.warning('Gemini OCR failed for PDF page %s/%s: %s', page_number, total_pages, exc)
        return ''


def extract_text_from_file(file_path: str) -> str:
    """Extract text from supported document, audio, and video files."""
    ext = os.path.splitext(file_path)[1].lower()
    try:
        if ext == '.pdf':
            text = _extract_pdf(file_path)
        elif ext in ('.docx', '.doc'):
            text = _extract_docx(file_path)
        elif ext == '.txt':
            text = _extract_txt(file_path)
        elif ext in ('.mp3', '.wav', '.m4a'):
            text = _extract_audio(file_path)
        elif ext in ('.mp4', '.mov'):
            text = _extract_video(file_path)
        else:
            logger.warning('Unsupported file type: %s', ext)
            text = ''
    except Exception as exc:
        logger.error('Text extraction failed for %s: %s', file_path, exc)
        text = ''

    cleaned = _clean_extracted_text(text)
    return cleaned or 'No extractable text found in the uploaded file.'


def extract_text_from_bytes(file_name: str, file_data: bytes) -> str:
    """Extract text from in-memory file bytes by using a temporary file."""
    if not file_data:
        return 'No extractable text found in the uploaded file.'

    ext = os.path.splitext(file_name or '')[1].lower() or '.bin'
    temp_path = None
    try:
        with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as temp_file:
            temp_file.write(file_data)
            temp_path = temp_file.name
        return extract_text_from_file(temp_path)
    except Exception as exc:
        logger.error('Text extraction failed for in-memory file %s: %s', file_name, exc)
        return 'No extractable text found in the uploaded file.'
    finally:
        if temp_path and os.path.exists(temp_path):
            os.remove(temp_path)


def extract_text_from_pdf(file_path: str) -> str:
    """Backward compatible alias for previous function name."""
    return extract_text_from_file(file_path)
